"""Build the local, offline data snapshot (data/snapshot/) from previously cached raw downloads.

BUILD TIME ONLY. This script never touches the network: it reads raw files that were fetched earlier
(after robots.txt / terms checks) from apis.datos.gob.ar and datos.gob.ar resources, verifies their
sha256, keeps only datasets whose licence metadata says Creative Commons Attribution 4.0, and writes
small CSVs + catalog.json. The runtime (src/apertus_qa) only ever reads data/snapshot/.

Usage:  RAW_CACHE_ROOT=/path/to/raw/caches python3 src/tools/build_snapshot.py
Layout expected under RAW_CACHE_ROOT (see docs/DATA.md):
  ipc/                 ipc_nivel.json (+ .meta.json), LICENCE_EVIDENCE.json, dataset_page.html
  b1/<slug>/           series-API JSON caches (poblacion, exportaciones, tipo_cambio, salarios_vs_precios)
  b2/<slug>/           series-API JSON / raw files + LICENSE.json (turismo, energia, petroleo, subte,
                       nacimientos, vuelos_anac, muertes_viales)
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]          # track_2b/
OUT = ROOT / "data" / "snapshot"
RAW = Path(os.environ.get("RAW_CACHE_ROOT", "")).expanduser()
CC_BY = {"Creative Commons Attribution 4.0", "CC-BY-4.0"}
API_TERMS = "https://datosgobar.github.io/series-tiempo-ar-api/terms/"


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def api_cache(path: Path) -> tuple[dict, dict]:
    j = json.loads(path.read_text())
    mp = path.with_name(path.stem + ".meta.json")
    m = json.loads(mp.read_text()) if mp.exists() else {}
    if m.get("sha256"):
        assert sha(path) == m["sha256"], f"{path} changed since fetch"
    return j, m


def series_from_api(j: dict, sid: str) -> tuple[list[tuple[str, float]], dict]:
    ids = [mm["field"]["id"] for mm in j["meta"][1:]]
    k = ids.index(sid) + 1
    meta = j["meta"][k]
    pts = [(r[0], float(r[k])) for r in j["data"] if r[k] is not None]
    return pts, meta


ENTRIES: list[dict] = []
FILES: dict[str, list[tuple[str, float]]] = {}
TABLES: dict[str, list[dict]] = {}


def add_series(sid, pts, *, title, topic, unit, kind, freq, source, dataset, licence, evidence, url, fetched_at,
               raw_sha, keywords, group=None, notes=(), forbidden=(), publisher=None, dataset_url=None,
               catalog_licence=None, sum_ok=False, scale=1.0):
    if licence not in CC_BY:
        print(f"SKIP {sid}: licence {licence!r} is not CC BY 4.0", file=sys.stderr)
        return
    pts = [(d[:10], v * scale) for d, v in pts]
    fname = "series/" + sid.replace("/", "_").replace("@", "__") + ".csv"
    FILES[fname] = pts
    ENTRIES.append({
        "id": sid, "title": title, "topic": topic, "unit": unit, "kind": kind, "frequency": freq,
        "aggregation_ok": sum_ok, "group": group, "keywords": list(keywords),
        "coverage": {"start": pts[0][0], "end": pts[-1][0], "points": len(pts)},
        "file": fname, "source": source, "publisher": publisher, "dataset": dataset, "dataset_url": dataset_url,
        "licence": "CC BY 4.0", "licence_raw": licence, "catalog_licence": catalog_licence,
        "licence_evidence": evidence, "api_terms": API_TERMS if "apis.datos.gob.ar" in (url or "") else None,
        "retrieved_from": url, "fetched_at": fetched_at, "raw_sha256": raw_sha,
        "notes": list(notes), "forbidden_words": list(forbidden),
    })


def add_table(tid, rows, **kw):
    if kw["licence"] not in CC_BY:
        print(f"SKIP {tid}: licence not CC BY 4.0", file=sys.stderr)
        return
    fname = f"tables/{tid}.csv"
    TABLES[fname] = rows
    ENTRIES.append({"id": tid, "kind": "table", "file": fname, "rows": len(rows), "licence": "CC BY 4.0",
                    "licence_raw": kw.pop("licence"), **kw})


def from_api(rel: str, sid: str, **kw):
    p = RAW / rel
    j, m = api_cache(p)
    pts, meta = series_from_api(j, sid)
    ds = meta["dataset"]
    kw.setdefault("source", ds.get("source") or (ds.get("publisher") or {}).get("name"))
    kw.setdefault("publisher", (ds.get("publisher") or {}).get("name"))
    kw.setdefault("dataset", ds.get("title"))
    add_series(sid, pts, licence=ds.get("license"), catalog_licence=(meta.get("catalog") or {}).get("license"),
               evidence=f"Series API metadata=full: dataset.license = {ds.get('license')!r}", url=m.get("url"),
               fetched_at=m.get("fetched_at"), raw_sha=sha(p), **kw)
    return pts, meta


def main():
    if not RAW.is_dir():
        sys.exit("Set RAW_CACHE_ROOT to the raw-cache directory (build time only; see docs/DATA.md)")
    # ---------- IPC (INDEC) ----------
    lic = json.loads((RAW / "ipc/LICENCE_EVIDENCE.json").read_text())
    assert lic["dataset_page"]["license"] in CC_BY and lic["api_metadata_full"]["dataset.license"] in CC_BY
    ipc, _ = from_api("ipc/ipc_nivel.json", "148.3_INIVELNAL_DICI_M_26",
                      title="IPC Nivel General Nacional (base dic-2016 = 100)", topic="precios", unit="índice (dic-2016=100)",
                      kind="index", freq="month", keywords=["ipc", "inflación", "inflacion", "precios", "índice de precios"],
                      dataset_url=lic["dataset_page"]["url"],
                      notes=["La inflación mensual e interanual se calcula a partir del índice (cociente − 1).",
                             "abr-2019: el valor de la API difiere levemente del CSV de INDEC (4e-4 relativo); se usa el de la API."])
    # IPC dic-dic annual % (derived at build time from the level series; cross-checked with an API-computed cache)
    lvl = dict(ipc)
    dd = [(f"{y}-01-01", lvl[f"{y}-12-01"] / lvl[f"{y-1}-12-01"] - 1) for y in range(2017, 2026)]
    chk, _ = api_cache(RAW / "b1/salarios_vs_precios/anual_dic_dic.json")
    api_dd = dict(series_from_api(chk, "148.3_INIVELNAL_DICI_M_26")[0])
    for d, v in dd:
        if d in api_dd:
            assert abs(api_dd[d] - v) < 1e-9, (d, v, api_dd[d])
    e = [x for x in ENTRIES if x["id"] == "148.3_INIVELNAL_DICI_M_26"][0]
    add_series("148.3_INIVELNAL_DICI_M_26@dic_dic", dd, title="Inflación anual diciembre a diciembre (IPC Nacional)",
               topic="precios", unit="variación % (proporción)", kind="pct", freq="year",
               keywords=["inflación anual", "inflacion anual", "diciembre a diciembre", "ipc"],
               source=e["source"], publisher=e["publisher"], dataset=e["dataset"], dataset_url=e["dataset_url"],
               licence="Creative Commons Attribution 4.0", evidence="derivada de 148.3_INIVELNAL_DICI_M_26 (CC BY 4.0)",
               url=e["retrieved_from"], fetched_at=e["fetched_at"], raw_sha=e["raw_sha256"], group="dic_dic",
               notes=["Derivada: índice dic(año) / índice dic(año−1) − 1. Coincide con la API (representation_mode=percent_change, collapse=year)."])
    # ---------- Salarios (INDEC), annual dic-dic % as published by the API ----------
    from_api("b1/salarios_vs_precios/anual_dic_dic.json", "149.1_TL_REGIADO_OCTU_0_16",
             title="Índice de salarios, empleo registrado: variación anual diciembre a diciembre", topic="salarios",
             unit="variación % (proporción)", kind="pct", freq="year", group="dic_dic",
             keywords=["salarios", "sueldos", "salario registrado", "índice de salarios"],
             notes=["Solo la variación dic–dic (API: collapse=year, end_of_period, percent_change); no el nivel del índice.",
                    "Cobertura en el snapshot: 2018–2025."])
    # ---------- Población, tipo de cambio (INDEC + SPE) ----------
    from_api("b1/poblacion/pob.json", "9.1_POB_2004_A_9", title="Población de Argentina", topic="población",
             unit="habitantes", kind="level", freq="year", keywords=["población", "poblacion", "habitantes"],
             notes=["Serie anual publicada junto al PBI en dólares (INDEC y Secretaría de Política Económica)."])
    from_api("b1/tipo_cambio/tc.json", "9.1_TU_2004_A_17", title="Tipo de cambio nominal peso/dólar (anual)",
             topic="tipo de cambio", unit="pesos por dólar", kind="level", freq="year",
             keywords=["tipo de cambio", "dólar", "dolar", "cotización"],
             notes=["Valor anual usado en la serie de PBI en dólares; no es la cotización diaria ni un dólar paralelo."])
    # ---------- Exportaciones (INDEC) ----------
    names = {"350.1_TOTAL_EXPO_PP__22": "Productos primarios", "350.1_TOTAL_EXPOMOA__23": "Manufacturas de origen agropecuario (MOA)",
             "350.1_TOTAL_EXPOMOI__23": "Manufacturas de origen industrial (MOI)", "350.1_TOTAL_EXPOCYE__23": "Combustibles y energía",
             "350.1_TOTAL_EXPONES__39": "Exportaciones totales"}
    for sid, nm in names.items():
        tot = sid.endswith("NES__39")
        from_api("b1/exportaciones/expo.json", sid, title=f"Exportaciones: {nm}", topic="comercio exterior",
                 unit="millones de dólares", kind="flow", freq="year", group=None if tot else "exportaciones_rubros",
                 keywords=["exportaciones", "exportó", "exporto", "ventas al exterior"] + ([nm.lower()] if not tot else ["total"]),
                 notes=["Dataset 'Exportaciones por provincia y rubro' (datos anuales); el total puede diferir del ICA publicado."])
    # ---------- Turismo emisivo (Subsecretaría de Turismo) ----------
    j, _ = api_cache(RAW / "b2/turismo_emisivo/te_turistas.json")
    for mm in j["meta"][1:]:
        sid = mm["field"]["id"]
        desc = mm["field"]["description"].replace("Estados Undicos", "Estados Unidos")
        via = "aerea" if "aérea" in desc else "terrestre" if "terrestre" in desc else "maritima_fluvial"
        from_api("b2/turismo_emisivo/te_turistas.json", sid, title=desc, topic="turismo", unit="viajes de turistas",
                 kind="flow", freq="month", group=f"turismo_emisivo_{via}", sum_ok=True,
                 source="Subsecretaría de Turismo, Dirección de Mercados y Estadísticas (datos de la Dirección Nacional de Migraciones)",
                 keywords=["turismo", "turistas", "viajes", "emisivo", "residentes", "viajaron"],
                 notes=["Licencia del dataset: CC BY 4.0 (metadato dataset.license). El catálogo de turismo declara ODbL-1.0 a nivel catálogo; rige la licencia del dataset."])
    # ---------- Energía (CAMMESA) ----------
    for sid, nm in [("367.1_POTENCIA_ILAR__24", "solar"), ("367.1_POTENCIA_IICA__25", "eólica"), ("367.1_POTENCIA_ITAL__24", "total")]:
        from_api("b2/energia_renovable/potencia_instalada.json", sid, title=f"Potencia instalada {nm} (fin de año)",
                 topic="energía", unit="MW", kind="level", freq="year", group="potencia_instalada" if nm != "total" else None,
                 keywords=["potencia instalada", "energía", "energia", nm, "renovable", "MW"])
    # ---------- Petróleo (Secretaría de Energía) ----------
    from_api("b2/petroleo/petroleo_crudo.json", "363.3_PRODUCCIONUDO__28", title="Producción de petróleo crudo",
             topic="energía", unit="miles de m³", kind="flow", freq="month", sum_ok=True,
             keywords=["petróleo", "petroleo", "crudo", "producción de petróleo"], forbidden=["récord", "record", "histórico"],
             notes=["No usar 'récord': la serie del snapshot empieza en 1996."])
    # ---------- Subte (INDEC) ----------
    xc = json.loads((RAW / "b2/subte_crosscheck.json").read_text())
    assert xc["max_abs_diff_thousands"] == 0.0
    from_api("b2/subte/subte_pasajeros.json", "302.3_TRANSP_PASSAJ_0_S_38", title="Pasajeros de subterráneo y premetro (CABA)",
             topic="transporte", unit="miles de pasajeros", kind="flow", freq="month", sum_ok=True,
             keywords=["subte", "subterráneo", "premetro", "pasajeros"],
             notes=["El metadato de unidades dice '2004=100', pero los valores son miles de pasajeros (verificado mes a mes contra el cuadro 9 del ISSP de INDEC, 174 meses, diferencia 0)."])
    # ---------- Nacimientos (DEIS, Ministerio de Salud), XLSX resource ----------
    lic_n = json.loads((RAW / "b2/nacimientos/LICENSE.json").read_text())
    page = lic_n["sources"][0]
    rawp = RAW / "b2/nacimientos/nacidos_vivos.csv"
    meta_n = json.loads((RAW / "b2/nacimientos/nacidos_vivos.meta.json").read_text())
    assert sha(rawp) == meta_n["sha256"]
    import openpyxl  # build-time only dependency
    wb = openpyxl.load_workbook(io.BytesIO(rawp.read_bytes()), read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    assert rows[0][0] == "anio" and rows[0][1] == "total_argentina"
    end_year = int(page["temporal_end"][:4])
    pts = [(f"{r[0].year}-01-01", float(r[1])) for r in rows[1:] if r[0] is not None and r[1] is not None and r[0].year <= end_year]
    add_series("deis_nacidos_vivos_total_pais", pts, title="Nacidos vivos, total del país", topic="población",
               unit="nacimientos", kind="flow", freq="year", source="Dirección de Estadísticas e Información de Salud (DEIS), Ministerio de Salud",
               dataset=page["title"].replace(" - Dataset", ""), dataset_url=page["url"], licence=page["license"],
               evidence="datos.gob.ar dataset page, campo 'Licencia'", url=meta_n["url"], fetched_at=meta_n["fetched_at"],
               raw_sha=meta_n["sha256"], keywords=["nacimientos", "nacidos vivos", "natalidad", "nacieron", "bebés"],
               notes=[f"Cobertura declarada del dataset: {page['temporal_start'][:4]}–{end_year}. Los datos llegan a {end_year}.",
                      "El archivo trae filas posteriores fuera de la cobertura declarada; no se usan.",
                      "Hay años sin dato en la serie histórica."])
    # ---------- Vuelos ANAC 2025 (counts only, never passengers) ----------
    lic_v = json.loads((RAW / "b2/vuelos_anac/LICENSE.json").read_text())
    meta_v = json.loads((RAW / "b2/vuelos_anac/anac_2025.meta.json").read_text())
    rawv = RAW / "b2/vuelos_anac/anac_2025.csv"
    assert sha(rawv) == meta_v["sha256"]
    ap, dep, years = Counter(), Counter(), Counter()
    with rawv.open(encoding="utf-8-sig", newline="") as f:
        rd = csv.reader(f, delimiter=";")
        h = next(rd)
        assert h[2].startswith("Clase de Vuelo") and h[5] == "Aeropuerto" and h[6] == "Origen / Destino"
        for r in rd:
            years[r[0][-4:]] += 1
            if r[2] != "Regular":
                continue
            ap[r[5]] += 1
            if r[3] == "Doméstico" and r[4] == "Despegue":
                dep[tuple(sorted((r[5], r[6])))] += 1
    assert set(years) == {"2025"}, years
    common = dict(topic="transporte", source="Administración Nacional de Aviación Civil (ANAC)",
                  dataset=lic_v["sources"][0]["title"].replace(" - Dataset", ""), dataset_url=lic_v["sources"][0]["url"],
                  licence=lic_v["sources"][0]["license"], licence_evidence="datos.gob.ar dataset page, campo 'Licencia'",
                  retrieved_from=meta_v["url"], fetched_at=meta_v["fetched_at"], raw_sha256=meta_v["sha256"], period="2025",
                  frequency="year")
    labels = {"AER": "Aeroparque", "EZE": "Ezeiza", "CBA": "Córdoba", "DOZ": "Mendoza", "BAR": "Bariloche", "IGU": "Iguazú",
              "SAL": "Salta", "NEU": "Neuquén", "USU": "Ushuaia", "TUC": "Tucumán", "ECA": "El Calafate",
              "CRV": "Comodoro Rivadavia"}
    common["labels"] = labels
    add_table("anac_2025_movimientos_regulares_por_aeropuerto",
              [{"clave": k, "valor": v} for k, v in ap.most_common()],
              title="Movimientos regulares (aterrizajes + despegues) por aeropuerto, 2025", unit="movimientos",
              keywords=["aeropuerto", "aeropuertos", "movimientos", "vuelos", "aterrizajes", "despegues"],
              notes=["Cuenta filas Clase de Vuelo = Regular. Solo conteos, nunca pasajeros.",
                     "Códigos de aeropuerto según el archivo de ANAC (p. ej. AER = Aeroparque, EZE = Ezeiza, CBA = Córdoba)."],
              forbidden_words=["pasajeros"], **common)
    add_table("anac_2025_vuelos_cabotaje_por_ruta",
              [{"clave": f"{a}-{b}", "valor": v} for (a, b), v in dep.most_common()],
              title="Vuelos regulares de cabotaje por ruta (ambos sentidos), 2025", unit="vuelos",
              keywords=["ruta", "rutas", "cabotaje", "vuelos", "aérea", "aerea"],
              notes=["Vuelos = despegues regulares domésticos entre los dos aeropuertos, sumando ambos sentidos.",
                     "Solo conteos de vuelos; las columnas de pasajeros de ANAC no se usan."],
              forbidden_words=["pasajeros"], **common)
    # ---------- Muertes viales (SNIC, Ministerio de Seguridad): aggregates only, never microdata ----------
    lic_m = json.loads((RAW / "b2/muertes_viales/LICENSE.json").read_text())
    meta_m = json.loads((RAW / "b2/muertes_viales/sat_mv_bu.meta.json").read_text())
    rawm = RAW / "b2/muertes_viales/sat_mv_bu.csv"
    assert sha(rawm) == meta_m["sha256"]
    per_year, veh = Counter(), Counter()
    with rawm.open(encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            if r["tipo_persona"] != "Víctima":
                continue
            per_year[r["anio"]] += 1
            if r["anio"] == "2024":
                veh[r["victima_vehiculo"] or "Sin dato"] += 1
    pg = lic_m["sources"][0]
    add_series("snic_sat_mv_victimas_anual", [(f"{y}-01-01", float(n)) for y, n in sorted(per_year.items())],
               title="Víctimas fatales de siniestros viales (Sistema de Alerta Temprana)", topic="seguridad vial",
               unit="víctimas", kind="count", freq="year", source="Ministerio de Seguridad, SNIC, Sistema de Alerta Temprana",
               dataset=pg["title"].replace(" - Dataset", ""), dataset_url=pg["url"], licence=pg["license"],
               evidence="datos.gob.ar dataset page, campo 'Licencia'", url=meta_m["url"], fetched_at=meta_m["fetched_at"],
               raw_sha=meta_m["sha256"], keywords=["siniestros viales", "muertes viales", "víctimas", "tránsito", "accidentes"],
               notes=["Agregado del snapshot: conteo de filas tipo_persona = Víctima por año (microdatos no incluidos).",
                      "Decir 'siniestros viales', no 'accidentes'. Puede diferir de otras fuentes (p. ej. ANSV)."],
               forbidden=["accidente", "accidentes"])
    add_table("snic_sat_mv_victimas_por_vehiculo_2024", [{"clave": k, "valor": v} for k, v in veh.most_common()],
              title="Víctimas fatales de siniestros viales por vehículo de la víctima, 2024", unit="víctimas",
              topic="seguridad vial", source="Ministerio de Seguridad, SNIC, Sistema de Alerta Temprana",
              dataset=pg["title"].replace(" - Dataset", ""), dataset_url=pg["url"], licence=pg["license"],
              licence_evidence="datos.gob.ar dataset page, campo 'Licencia'", retrieved_from=meta_m["url"],
              fetched_at=meta_m["fetched_at"], raw_sha256=meta_m["sha256"], period="2024", frequency="year",
              keywords=["moto", "motocicleta", "vehículo", "auto", "peatón", "siniestros viales"],
              notes=["Conteo por victima_vehiculo. 'Motocicleta' no incluye 'Ciclomotor'."], forbidden_words=["accidente", "accidentes"])

    # ---------- write ----------
    if OUT.exists():
        for p in sorted(OUT.rglob("*"), reverse=True):
            p.unlink() if p.is_file() else p.rmdir()
    (OUT / "series").mkdir(parents=True)
    (OUT / "tables").mkdir(parents=True)
    hashes = {}
    for fname, pts in FILES.items():
        s = "fecha,valor\n" + "".join(f"{d},{repr(v)}\n" for d, v in pts)
        (OUT / fname).write_text(s)
        hashes[fname] = hashlib.sha256(s.encode()).hexdigest()
    for fname, rows in TABLES.items():
        s = "clave,valor\n" + "".join(f"{r['clave']},{r['valor']}\n" for r in rows)
        (OUT / fname).write_text(s)
        hashes[fname] = hashlib.sha256(s.encode()).hexdigest()
    for e in ENTRIES:
        e["sha256"] = hashes[e["file"]]
    cat = {"snapshot_version": 1, "built_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
           "built_from": "raw caches fetched 2026-10-01/02 from apis.datos.gob.ar and datos.gob.ar resources (robots.txt checked)",
           "licence_policy": "Only datasets whose licence metadata is Creative Commons Attribution 4.0 are included.",
           "attribution": "Fuentes: INDEC, Ministerio de Economía, Secretaría de Energía, CAMMESA, Subsecretaría de Turismo, "
                          "Ministerio de Salud (DEIS), ANAC, Ministerio de Seguridad (SNIC), vía datos.gob.ar / apis.datos.gob.ar. "
                          "Licencia CC BY 4.0. Datos transformados (agregados y variaciones calculadas) por este proyecto.",
           "entries": ENTRIES}
    (OUT / "catalog.json").write_text(json.dumps(cat, ensure_ascii=False, indent=1))
    print(f"snapshot: {len(ENTRIES)} entries, {sum(1 for e in ENTRIES if e['kind'] != 'table')} series, "
          f"{len(TABLES)} tables -> {OUT}")


if __name__ == "__main__":
    main()
